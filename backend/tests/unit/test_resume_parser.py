import io

import pytest

from ai.parsers.resume_parser import parse_resume_file
from app.core.exceptions import FileTooLargeError, UnsupportedFileTypeError, ValidationFailedError


def test_docx_extraction_returns_text():
    import docx

    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Python Developer with 5 years experience.")
    document.save(buf)

    text = parse_resume_file(buf.getvalue(), "resume.docx")
    assert "Jane Doe" in text
    assert "Python Developer" in text


def test_pdf_extraction_returns_text():
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Jane Doe\nPython Developer with 5 years experience.")
    pdf_bytes = doc.tobytes()
    doc.close()

    text = parse_resume_file(pdf_bytes, "resume.pdf")
    assert "Jane Doe" in text


def test_empty_pdf_raises_error():
    import fitz

    doc = fitz.open()
    doc.new_page()
    pdf_bytes = doc.tobytes()
    doc.close()

    with pytest.raises(ValidationFailedError):
        parse_resume_file(pdf_bytes, "empty.pdf")


def test_oversized_file_raises_error():
    huge = b"x" * (11 * 1024 * 1024)  # 11MB > default 10MB limit
    with pytest.raises(FileTooLargeError):
        parse_resume_file(huge, "resume.txt")


def test_unsupported_type_raises_error():
    with pytest.raises(UnsupportedFileTypeError):
        parse_resume_file(b"hello", "resume.exe")


def test_text_cleaned_correctly():
    raw = b"Jane   Doe\r\n\r\n\r\n\r\nSkills:\ndevelop-\nment"
    text = parse_resume_file(raw, "resume.txt")
    assert "\r" not in text
    assert "\n\n\n" not in text
    assert "development" in text
