"""
ai/parsers/resume_parser.py

PDF (PyMuPDF) / DOCX (python-docx) / TXT -> raw text extraction.
Deterministic, no LLM involved (that's ai/extractors/resume_extractor.py).

Enforces: max file size (MAX_FILE_SIZE_MB), supported extensions,
non-empty extracted text. Fails loudly (Golden Rule 9) rather than
returning empty/garbage text silently.
"""
from __future__ import annotations

from app.core.config import get_settings
from app.core.exceptions import FileTooLargeError, UnsupportedFileTypeError, ValidationFailedError

from .text_cleaner import clean_text

SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}


def _extract_pdf(file_bytes: bytes) -> str:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=file_bytes, filetype="pdf")
    try:
        pages = [page.get_text() for page in doc]
    finally:
        doc.close()
    return "\n".join(pages)


def _extract_docx(file_bytes: bytes) -> str:
    import io

    import docx

    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValidationFailedError("Could not decode .txt file with any supported encoding.")


def parse_resume_file(file_bytes: bytes, file_name: str) -> str:
    """
    Extract clean, normalized text from an uploaded resume file.
    Raises FileTooLargeError, UnsupportedFileTypeError, or
    ValidationFailedError (empty result) as appropriate.
    """
    settings = get_settings()

    max_bytes = settings.max_file_size_mb * 1024 * 1024
    if len(file_bytes) > max_bytes:
        raise FileTooLargeError(
            f"File exceeds the {settings.max_file_size_mb}MB limit.",
            details={"file_name": file_name, "size_bytes": len(file_bytes)},
        )

    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type: .{ext}. Supported: {sorted(SUPPORTED_EXTENSIONS)}.",
            details={"file_name": file_name, "extension": ext},
        )

    extractors = {"pdf": _extract_pdf, "docx": _extract_docx, "txt": _extract_txt}
    raw_text = extractors[ext](file_bytes)
    cleaned = clean_text(raw_text)

    if not cleaned:
        raise ValidationFailedError(
            "No extractable text found in resume file (empty or scanned-image PDF?).",
            details={"file_name": file_name},
        )

    return cleaned
